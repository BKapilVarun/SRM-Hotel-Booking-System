import urllib.request
import json
from openpyxl import load_workbook

# 1. Test Catalog
res = urllib.request.urlopen('http://127.0.0.1:5000/api/catalog')
cat = json.loads(res.read())
print(f"Catalog loaded: {len(cat['rooms'])} rooms, {len(cat['menu'])} dishes")

# 2. Test Room Booking
room_payload = json.dumps({
    'guest_name': 'Lord Sterling Cross',
    'email': 'sterling@cross-estates.com',
    'phone': '+1 (555) 789-9900',
    'room_type': 'The Presidential Sky Penthouse',
    'price_per_night': 850,
    'check_in': '2026-10-15',
    'check_out': '2026-10-18',
    'nights': 3,
    'guests': 2,
    'total_amount': 2550,
    'payment_method': 'Platinum Card Guarantee',
    'special_requests': 'Chilled vintage champagne upon arrival'
}).encode('utf-8')

req = urllib.request.Request('http://127.0.0.1:5000/api/book-room', data=room_payload, headers={'Content-Type': 'application/json'})
res = urllib.request.urlopen(req)
room_res = json.loads(res.read())
print('Room Booking Result:', room_res['success'], room_res['receipt']['booking_id'])

# 3. Test Food Order
food_payload = json.dumps({
    'customer_name': 'Lady Vivian Vance',
    'phone': '+1 (555) 432-1100',
    'email': 'vivian@vance.org',
    'service_type': 'Room Service',
    'destination_number': 'Villa 04',
    'preferred_time': 'Evening Sunset (7:30 PM)',
    'dietary_notes': 'No peanuts, extra truffle',
    'subtotal': 142.0,
    'tax_service': 14.2,
    'grand_total': 156.2,
    'payment_method': 'Charge to Suite Folio',
    'items': [
        {'name': 'A5 Wagyu Striploin & Black Truffle', 'qty': 1, 'price': 78},
        {'name': 'Brittany Blue Lobster Risotto', 'qty': 1, 'price': 64}
    ]
}).encode('utf-8')

req2 = urllib.request.Request('http://127.0.0.1:5000/api/order-food', data=food_payload, headers={'Content-Type': 'application/json'})
res2 = urllib.request.urlopen(req2)
food_res = json.loads(res2.read())
print('Food Order Result:', food_res['success'], food_res['receipt']['order_id'])

# 4. Test Unified Booking (2-in-1)
unified_payload = json.dumps({
    'room': {
        'guest_name': 'Countess Eleanor Dupont',
        'email': 'eleanor@dupont.fr',
        'phone': '+33 6 12 34 56 78',
        'room_type': 'Emerald Oceanfront Pool Villa',
        'price_per_night': 620,
        'check_in': '2026-11-01',
        'check_out': '2026-11-05',
        'nights': 4,
        'guests': 2,
        'total_amount': 2480,
        'payment_method': 'Stay & Dine Pass Guarantee'
    },
    'food': {
        'customer_name': 'Countess Eleanor Dupont',
        'phone': '+33 6 12 34 56 78',
        'email': 'eleanor@dupont.fr',
        'items': [
            {'name': 'The Royal Aegean Sunrise Breakfast', 'qty': 2, 'price': 38},
            {'name': 'Vintage Dom Pérignon Brut', 'qty': 1, 'price': 45}
        ],
        'preferred_time': 'Every Morning at 8:30 AM',
        'subtotal': 121.0,
        'tax_service': 12.1,
        'grand_total': 133.1
    }
}).encode('utf-8')

req3 = urllib.request.Request('http://127.0.0.1:5000/api/book-unified', data=unified_payload, headers={'Content-Type': 'application/json'})
res3 = urllib.request.urlopen(req3)
unified_res = json.loads(res3.read())
print('Unified Booking Result:', unified_res['success'], 'Room:', unified_res['receipt']['room_booking_id'], 'Food:', unified_res['receipt']['food_order_id'])

# 5. Verify Excel Files
wb_room = load_workbook('excel_data/hotel_bookings.xlsx')
ws_room = wb_room.active
print(f'\nhotel_bookings.xlsx total rows: {ws_room.max_row}')
for r in range(1, ws_room.max_row + 1):
    vals = [ws_room.cell(r, c).value for c in range(1, 8)]
    print(f'  Row {r}: {vals}')

wb_food = load_workbook('excel_data/food_orders.xlsx')
ws_food = wb_food.active
print(f'\nfood_orders.xlsx total rows: {ws_food.max_row}')
for r in range(1, ws_food.max_row + 1):
    vals = [ws_food.cell(r, c).value for c in range(1, 8)]
    print(f'  Row {r}: {vals}')
